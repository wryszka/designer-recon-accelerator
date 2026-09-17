# Databricks notebook source
# MAGIC %md
# MAGIC # UC1 · the receipt behind the Designer flow (NOT shown in the room)
# MAGIC
# MAGIC The **coded mirror** of the no-code **Lakeflow Designer** flow — not what a customer builds or
# MAGIC maintains. In the room the **Designer canvas** does the reconciliation (no code) and the analyst gets
# MAGIC the Excel from a **standard download**. Here we prove it ties out.
# MAGIC 1. LEFT-join the rolling file (`cf_prior_rec`) to this period's figures (`cf_period_extract`) on
# MAGIC    **Account No. + Category**,
# MAGIC 2. append the **`Period NN` + `Period NN Control`** pair,
# MAGIC 3. derive **`Status`** on each account's **Closing Balance per Bank** row (Reconciled when its Control is 0),
# MAGIC 4. **demo-only QA:** check it matches the synthetic benchmark to the penny (no benchmark exists in a real
# MAGIC    deployment — scaffolding, never shown),
# MAGIC 5. **formatted export (hidden):** a small reusable template drops `.csv` + `.xlsx` into the folder.

# COMMAND ----------

# MAGIC %pip install openpyxl -q
# MAGIC dbutils.library.restartPython()

# COMMAND ----------

dbutils.widgets.text("catalog_name", "lr_dev_aws_us_catalog")
dbutils.widgets.text("schema_name", "designer_recon_demo")
dbutils.widgets.text("landing_volume_name", "recon_landing")
catalog = dbutils.widgets.get("catalog_name"); schema = dbutils.widgets.get("schema_name")
volume = dbutils.widgets.get("landing_volume_name")
fqn = f"{catalog}.{schema}"; vroot = f"/Volumes/{catalog}/{schema}/{volume}/uc1"

import os, shutil, tempfile
import numpy as np, pandas as pd

# COMMAND ----------

# ---- 1-3. join -> append period pair -> Status (this is what the Designer canvas produces) ----
prior = spark.table(f"{fqn}.cf_prior_rec").toPandas()
extract = spark.table(f"{fqn}.cf_period_extract").toPandas()

prior_nums = [int(c.split()[1]) for c in prior.columns if c.startswith("Period ") and not c.endswith("Control")]
CUR_P = max(prior_nums) + 1
PCOL, CCOL = f"Period {CUR_P:02d}", f"Period {CUR_P:02d} Control"

out = prior.merge(extract[["Account No.", "Category", "current_period", "period_control"]],
                  on=["Account No.", "Category"], how="left")
out = out.rename(columns={"current_period": PCOL, "period_control": CCOL})
out["Status"] = np.where(out["Category"] == "Closing Balance per Bank",
                         np.where(out[CCOL].abs() < 0.01, "Reconciled", "Exception"), None)

(spark.createDataFrame(out.astype(object).where(pd.notnull(out), None)).write.mode("overwrite")
    .option("overwriteSchema", "true").option("delta.columnMapping.mode", "name").saveAsTable(f"{fqn}.cf_cashflow_rec"))
spark.sql(f"COMMENT ON TABLE {fqn}.cf_cashflow_rec IS 'UC1 produced rolling rec (long): prior + this period Period NN / Control + per-account Status. Synthetic.'")

# COMMAND ----------

# ---- 4. DEMO-ONLY QA: parity vs the synthetic benchmark (no benchmark exists in production) ----
bench = spark.table(f"{fqn}.cf_benchmark").toPandas()
key = ["Account No.", "Category"]
a = out.sort_values(key).reset_index(drop=True)
b = bench.sort_values(key).reset_index(drop=True)
num = [c for c in a.columns if c.startswith("Period ")]
val_mism = int((a[num].fillna(0).round(2) - b[num].fillna(0).round(2)).abs().gt(0.005).sum().sum())
status_mism = int((a["Status"].fillna("") != b["Status"].fillna("")).sum())
print(f"PARITY: {'✅ to the penny' if val_mism == 0 and status_mism == 0 else f'❌ {val_mism} value / {status_mism} status diffs'}")
assert val_mism == 0 and status_mism == 0, "parity failed"

# COMMAND ----------

# ---- 5. formatted Excel out — a small reusable export template (HIDDEN setup; never shown in the room) ----
from openpyxl.styles import Font, PatternFill
os.makedirs(f"{vroot}/output", exist_ok=True)
out.to_csv(f"{vroot}/output/CashFlowRec_period{CUR_P:02d}.csv", index=False)
_t = tempfile.mkdtemp(); xpath = f"{_t}/o.xlsx"
exc_accts = set(out.loc[out["Status"] == "Exception", "Account No."])
with pd.ExcelWriter(xpath, engine="openpyxl") as xw:
    out.to_excel(xw, index=False, sheet_name="CashFlowRec")
    ws = xw.sheets["CashFlowRec"]
    acc_col = list(out.columns).index("Account No.") + 1
    num_cols = [i + 1 for i, c in enumerate(out.columns) if c.startswith("Period ")]
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1B3A4B")
    for r in range(2, ws.max_row + 1):
        for i in num_cols:
            ws.cell(r, i).number_format = "#,##0.00"
        if ws.cell(r, acc_col).value in exc_accts:                 # flag the exception account's lines in red
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).font = Font(color="C00000", bold=True)
shutil.copy(xpath, f"{vroot}/output/CashFlowRec_period{CUR_P:02d}.xlsx")

recon = int((out["Status"] == "Reconciled").sum()); exc = int((out["Status"] == "Exception").sum())
print(f"CSV + formatted Excel → {vroot}/output/  ·  {recon} reconciled, {exc} exception(s)")
display(spark.table(f"{fqn}.cf_cashflow_rec"))
