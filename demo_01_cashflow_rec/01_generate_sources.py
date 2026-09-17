# Databricks notebook source
# MAGIC %md
# MAGIC # UC1 · Generate the synthetic sources (real cash-flow-rec shape, synthetic values)
# MAGIC
# MAGIC Shapes mirror a real monthly cash-flow reconciliation; **all values are synthetic — no real names,
# MAGIC accounts, sort codes or PII**. 6 accounts kept legible.
# MAGIC
# MAGIC Lands in the `recon_landing` Volume under `uc1/`:
# MAGIC - `bank_recs/` — per-account **rec workbooks** (`.xlsx`, a `Header` sheet: account details + the 8
# MAGIC   reconciliation figures — opening/closing O/S items and balances per SAP and per Bank)
# MAGIC - `fallback/` — the account with no workbook → a **SAP balances** `.xlsx` + a **bank-statement** `.csv`
# MAGIC - `rolling/` — last period's **rolling file** (long format: one row per account × reconciliation line,
# MAGIC   a `Period NN` + `Period NN Control` pair appended each period)
# MAGIC
# MAGIC Tables: `cf_accounts`, `cf_prior_rec` (rolling baseline), `cf_benchmark` (parity oracle).
# MAGIC This period's `cf_period_extract` is produced by the **Autoloader ingest** (`00_`).

# COMMAND ----------

# MAGIC %pip install openpyxl -q
# MAGIC dbutils.library.restartPython()

# COMMAND ----------

dbutils.widgets.text("catalog_name", "lr_dev_aws_us_catalog")
dbutils.widgets.text("schema_name", "designer_recon_demo")
dbutils.widgets.text("landing_volume_name", "recon_landing")
dbutils.widgets.text("period", "2026-07")
dbutils.widgets.text("fy_start_month", "4")
dbutils.widgets.text("seed", "71")
catalog = dbutils.widgets.get("catalog_name"); schema = dbutils.widgets.get("schema_name")
volume = dbutils.widgets.get("landing_volume_name"); period = dbutils.widgets.get("period")
fy_start_month = int(dbutils.widgets.get("fy_start_month")); seed = int(dbutils.widgets.get("seed"))
fqn = f"{catalog}.{schema}"
spark.sql(f"CREATE SCHEMA IF NOT EXISTS {fqn}")
spark.sql(f"CREATE VOLUME IF NOT EXISTS {fqn}.{volume}")

# COMMAND ----------

import os, shutil, tempfile
import numpy as np, pandas as pd
from openpyxl import Workbook

TMP = tempfile.mkdtemp(); rng = np.random.default_rng(seed)
cur = pd.Period(period, freq="M")
start_year = cur.year if cur.month >= fy_start_month else cur.year - 1
n_periods = (cur - pd.Period(f"{start_year}-{fy_start_month:02d}", freq="M")).n + 1   # Apr..Jul -> 4
PNUM = lambda i: f"Period {i:02d}"                                                     # Period 01 .. Period NN
CUR_P = n_periods                                                                      # current period number

vroot = f"/Volumes/{catalog}/{schema}/{volume}/uc1"
shutil.rmtree(vroot, ignore_errors=True)
for sub in ["bank_recs", "fallback", "rolling", "output"]:
    os.makedirs(f"{vroot}/{sub}", exist_ok=True)

# the 8 reconciliation lines (one row each, per account) — Control lives only on the "Balance per Bank" rows
CATS = ["Opening O/S items on SAP not on bank", "Opening O/S items on bank not on SAP",
        "Opening Balance per SAP", "Opening Balance per Bank",
        "Closing O/S items on SAP not on Bank", "Closing O/S items on bank not on SAP",
        "Closing Balance per SAP", "Closing Balance per Bank"]
BANK_ROWS = {"Opening Balance per Bank", "Closing Balance per Bank"}

# COMMAND ----------

# ---- 6 synthetic accounts: 5 with workbooks (1 an exception), 1 missing -> fallback ----
ACCTS = [
    {"Account No.": "0010001", "Contact": "Analyst 1", "A/c Name": "Operating Account 001",  "GL A/c": "5700000001", "Hse Bank": "HBK01", "A/c ID": "AC100", "CUR": "GBP"},
    {"Account No.": "0010002", "Contact": "Analyst 1", "A/c Name": "Claims Settlement Account","GL A/c": "5700000002", "Hse Bank": "HBK01", "A/c ID": "AC200", "CUR": "GBP"},
    {"Account No.": "0010003", "Contact": "Analyst 2", "A/c Name": "Premium Receipts Account", "GL A/c": "5700000003", "Hse Bank": "HBK02", "A/c ID": "AC300", "CUR": "GBP"},  # exception
    {"Account No.": "0010004", "Contact": "Analyst 2", "A/c Name": "Reinsurance Account",      "GL A/c": "5700000004", "Hse Bank": "HBK02", "A/c ID": "AC400", "CUR": "GBP"},
    {"Account No.": "0010005", "Contact": "Analyst 3", "A/c Name": "Suspense Account",         "GL A/c": "5700000005", "Hse Bank": "HBK03", "A/c ID": "AC500", "CUR": "GBP"},
    {"Account No.": "0010006", "Contact": "Analyst 3", "A/c Name": "Broker Receipts Account",  "GL A/c": "5700000006", "Hse Bank": "HBK03", "A/c ID": "AC900", "CUR": "GBP"},  # missing -> fallback
]
EXCEPTION_I, MISSING_I = 2, 5

def figures(level, exception=False, no_os=False):
    """One period's 8 figures for an account. Bank = SAP + (bank-not-SAP O/S) - (SAP-not-bank O/S) => Control 0.
    no_os (the fallback account) => no O/S items, so Bank == SAP and the 2-cell fallback check still ties."""
    z = (lambda: 0.0) if no_os else (lambda: round(rng.uniform(0, 800), 2))
    o_sap = round(level, 2); o_sn = z(); o_bn = z()
    o_bank = round(o_sap + o_bn - o_sn, 2)
    c_sap = round(o_sap + rng.normal(0, 50_000), 2); c_sn = z(); c_bn = z()
    c_bank = round(c_sap + c_bn - c_sn, 2); c_ctl = 0.0
    if exception:
        c_bank = round(c_bank + 420.50, 2); c_ctl = round(c_bank - (c_sap + c_bn - c_sn), 2)   # breaks to +420.50
    vals = {CATS[0]: o_sn, CATS[1]: o_bn, CATS[2]: o_sap, CATS[3]: o_bank,
            CATS[4]: c_sn, CATS[5]: c_bn, CATS[6]: c_sap, CATS[7]: c_bank}
    ctl = {"Opening Balance per Bank": 0.0, "Closing Balance per Bank": c_ctl}
    return vals, ctl, c_sap

# per account: figures for every period (history reconciled; exception only bites in the current period)
levels = np.round(rng.uniform(150_000, 3_000_000, len(ACCTS)), 2)
per_acct = []   # per_acct[i] = {period_num: (vals, ctl)}
for i in range(len(ACCTS)):
    lvl = levels[i]; by_p = {}
    for p in range(1, n_periods + 1):
        is_exc = (i == EXCEPTION_I and p == CUR_P)
        vals, ctl, lvl = figures(lvl, exception=is_exc, no_os=(i == MISSING_I)); by_p[p] = (vals, ctl)
    per_acct.append(by_p)

# COMMAND ----------

# ---- rolling file (long): one row per account x reconciliation line ----
FB_CELLS = {"Closing Balance per SAP", "Closing Balance per Bank"}   # the 2 cells the fallback provides

def build_long(period_range, current_filter=False):
    rows = []
    for i, a in enumerate(ACCTS):
        for cat in CATS:
            row = dict(a); row["Category"] = cat
            for p in period_range:
                vals, ctl = per_acct[i][p]
                fb = current_filter and p == CUR_P and i == MISSING_I    # fallback: only 2 cells this period
                if fb and cat not in FB_CELLS:
                    row[PNUM(p)] = None; row[f"{PNUM(p)} Control"] = None; continue
                row[PNUM(p)] = vals[cat]
                if cat == "Closing Balance per Bank" and fb:
                    row[f"{PNUM(p)} Control"] = round(vals["Closing Balance per Bank"] - vals["Closing Balance per SAP"], 2)
                elif cat in BANK_ROWS:
                    row[f"{PNUM(p)} Control"] = ctl.get(cat)
                else:
                    row[f"{PNUM(p)} Control"] = None
            rows.append(row)
    return pd.DataFrame(rows)

prior = build_long(range(1, CUR_P))                # rolling baseline the analyst carries forward (full history)
prior_label = f"{cur.year}-{(cur.month-1):02d}"    # e.g. 2026-06

_x = f"{TMP}/roll.xlsx"; prior.to_excel(_x, index=False, sheet_name="Sheet1")
shutil.copy(_x, f"{vroot}/rolling/CashFlowRec_{prior_label}.xlsx")
prior.to_csv(f"{vroot}/rolling/CashFlowRec_{prior_label}.csv", index=False)

# COMMAND ----------

# ---- per-account rec workbooks: a Header sheet (account details + the 8 current-period figures) ----
def write_workbook(a, vals, ctl):
    wb = Workbook(); ws = wb.active; ws.title = "Header"
    ws["A3"] = "MONTH END RECONCILIATION"
    ws["A6"] = "Account Details"
    for r, (lab, key) in enumerate([("Company Code", "GL A/c"), ("GL Number", "GL A/c"),
                                     ("Housebank/Account ID", None), ("Account Number", "Account No."),
                                     ("GL Account Name", "A/c Name"), ("Account Currency", "CUR")], 8):
        ws[f"A{r}"] = lab
        ws[f"B{r}"] = f'{a["Hse Bank"]} {a["A/c ID"]}' if lab == "Housebank/Account ID" else a[key]
    for r, cat in enumerate(CATS, 18):                 # the 8 figures the ingest reads
        ws[f"A{r}"] = cat; ws[f"B{r}"] = float(vals[cat])
        if cat in BANK_ROWS: ws[f"C{r}"] = "Control"; ws[f"D{r}"] = float(ctl.get(cat, 0.0))
    fn = f'{a["GL A/c"]} {a["Hse Bank"]} {a["A/c ID"]} {a["A/c Name"]}.xlsx'
    tmp = f"{TMP}/wb.xlsx"; wb.save(tmp); shutil.copy(tmp, f"{vroot}/bank_recs/{fn}")

for i, a in enumerate(ACCTS):
    vals, ctl = per_acct[i][CUR_P]
    if i == MISSING_I:
        # fallback: SAP balances xlsx (Closing Balance per SAP) + bank-statement csv (Closing Balance per Bank)
        sap = pd.DataFrame([{"Currency": a["CUR"], "House Bank New": a["Hse Bank"], "Account ID New": a["A/c ID"],
                             "Amt Cum.lc.cur": vals["Closing Balance per SAP"], "Currency 2": a["CUR"]}])
        _s = f"{TMP}/sap.xlsx"; sap.to_excel(_s, index=False, sheet_name="Sheet1")
        shutil.copy(_s, f"{vroot}/fallback/SAP Bank balances {period}.xlsx")
        with open(f"{vroot}/fallback/Bank Statement balances {period}.csv", "w") as fh:
            fh.write("Bank statement export,,,,,,,,\n")
            fh.write("GRP,ACC ID,ACCT NO,TYPE,BANK CODE,CURR,BAL DT,AS AT,STMT BAL\n")
            fh.write(f',,{a["Account No."]},,00-00-00,{a["CUR"]},{cur.strftime("%d/%m/%Y")},Close,{vals["Closing Balance per Bank"]}\n')
    else:
        write_workbook(a, vals, ctl)

# COMMAND ----------

# ---- baseline + oracle tables ----
bench_rows = build_long(range(1, CUR_P + 1), current_filter=True)   # + current period (fallback = 2 cells)
# Status: derived per account on the Closing Balance per Bank row (Reconciled when its Control is 0)
cclab = f"{PNUM(CUR_P)} Control"
bench_rows["Status"] = np.where((bench_rows["Category"] == "Closing Balance per Bank")
                                & (bench_rows[cclab].abs() < 0.01), "Reconciled",
                        np.where(bench_rows["Category"] == "Closing Balance per Bank", "Exception", None))
accounts = pd.DataFrame(ACCTS)

# old cf_* tables predate column mapping — drop so they recreate with real headers
for t in ["cf_accounts", "cf_prior_rec", "cf_benchmark", "cf_period_extract", "cf_cashflow_rec"]:
    spark.sql(f"DROP TABLE IF EXISTS {fqn}.{t}")

def write(df, name, comment):
    (spark.createDataFrame(df.astype(object).where(pd.notnull(df), None)).write.mode("overwrite")
        .option("overwriteSchema", "true").option("delta.columnMapping.mode", "name")   # real headers ("Account No." etc.)
        .saveAsTable(f"{fqn}.{name}"))
    spark.sql(f"COMMENT ON TABLE {fqn}.{name} IS '{comment}'")

write(accounts, "cf_accounts", "UC1 account reference (synthetic).")
write(prior, "cf_prior_rec", "UC1 rolling cash-flow file (long: account x reconciliation line; a Period NN + Control pair per period). Synthetic.")
write(bench_rows, "cf_benchmark", "UC1 parity oracle: rolling + current Period pair + per-account Status on the Closing Balance per Bank row. Synthetic.")

n_exc = int((bench_rows["Status"] == "Exception").sum())
print(f"{len(ACCTS)} accounts (5 workbooks + 1 fallback) · exception on {ACCTS[EXCEPTION_I]['Account No.']} · current = {PNUM(CUR_P)}")
print("landed:", {s: os.listdir(f'{vroot}/{s}') for s in ['bank_recs', 'fallback', 'rolling']})
display(spark.table(f"{fqn}.cf_benchmark"))
